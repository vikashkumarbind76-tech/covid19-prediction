import os
import pickle
import csv
from datetime import datetime
import json
from flask import Flask, render_template, request, jsonify, redirect, url_for, session, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from pymongo import MongoClient, ReturnDocument
from pymongo.errors import DuplicateKeyError

# Document Generation Imports
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'covid19_prediction_system_secret_key_123!@#')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# MongoDB connection
MONGO_URI = os.environ.get('MONGO_URI', 'mongodb://localhost:27017/')
db = None
try:
    client = MongoClient(MONGO_URI, serverSelectionTimeoutMS=2000)
    # Check connection
    client.server_info()
    db = client['covid_prediction_system']
    print("Successfully connected to MongoDB.")
except Exception as e:
    print(f"Error connecting to MongoDB: {e}")
    db = None

# Define export paths
if os.environ.get('VERCEL') or os.environ.get('AWS_LAMBDA_FUNCTION_NAME'):
    EXPORTS_DIR = '/tmp/exports'
else:
    EXPORTS_DIR = os.path.join(BASE_DIR, 'exports')

# Ensure directories exist
try:
    os.makedirs(EXPORTS_DIR, exist_ok=True)
except Exception as e:
    print(f"Warning: Could not create exports directory {EXPORTS_DIR}: {e}")

# Admin Credentials
ADMIN_USERNAME = 'admin'
ADMIN_PASSWORD = 'admin123'

# Load the ML model
MODEL_PATH = os.path.join(BASE_DIR, 'covid19Model.pkl')
try:
    with open(MODEL_PATH, 'rb') as f:
        model = pickle.load(f)
    print("Machine learning model loaded successfully.")
except Exception as e:
    print(f"Error loading model from {MODEL_PATH}: {e}")
    model = None

# MongoDB helper functions
def is_db_connected():
    if db is not None:
        try:
            client.server_info()
            return True
        except Exception:
            return False
    return False

def get_next_sequence_value(sequence_name):
    if db is None:
        raise Exception("Database not connected.")
    result = db.counters.find_one_and_update(
        {'_id': sequence_name},
        {'$inc': {'sequence_value': 1}},
        upsert=True,
        return_document=ReturnDocument.AFTER
    )
    return result['sequence_value']

def db_get_user_by_username(username):
    if db is None:
        return None
    return db.users.find_one({'username': username.strip().lower()})

def db_get_user_by_id(user_id):
    if db is None:
        return None
    return db.users.find_one({'id': int(user_id)})

def db_create_user(user_id, user_data):
    if db is None:
        raise Exception("Database not connected.")
    db.users.insert_one(user_data)

def db_get_predictions_by_user_id(user_id):
    if db is None:
        return []
    records = list(db.predictions.find({'user_id': int(user_id), 'is_deleted': {'$ne': True}}).sort('timestamp', -1))
    return records

def db_create_prediction(pred_id, pred_data):
    if db is None:
        raise Exception("Database not connected.")
    db.predictions.insert_one(pred_data)

def db_get_prediction_by_id(pred_id):
    if db is None:
        return None
    return db.predictions.find_one({'id': int(pred_id), 'is_deleted': {'$ne': True}})

def db_get_all_predictions():
    if db is None:
        return []
    records = list(db.predictions.find({'is_deleted': {'$ne': True}}).sort('timestamp', -1))
    return records

def db_get_users_map():
    if db is None:
        return {}
    users = db.users.find({}, {'id': 1, 'username': 1})
    return {u['id']: u['username'] for u in users}

def db_delete_prediction(pred_id):
    if db is None:
        raise Exception("Database not connected.")
    # Soft delete: update is_deleted flag to permanently store the record
    db.predictions.update_one({'id': int(pred_id)}, {'$set': {'is_deleted': True}})

def init_db():
    if db is None:
        print("Database not connected. Skipping initialization.")
        return
    # Ensure unique index on username
    db.users.create_index('username', unique=True)
    # Ensure index on predictions id and user_id for faster queries
    db.predictions.create_index('id')
    db.predictions.create_index('user_id')
    print("MongoDB Database connected and initialized with indexes.")

init_db()


# ==========================================================================
# PUBLIC PORTAL / USER DASHBOARD ROUTES
# ==========================================================================

@app.route('/')
def index():
    if not session.get('user_id'):
        return redirect(url_for('login'))
        
    try:
        user_records = db_get_predictions_by_user_id(session['user_id'])
    except Exception as e:
        print(f"Error fetching user records: {e}")
        user_records = []
    
    # Compile totals for the dashboard
    total = len(user_records)
    positives = sum(1 for r in user_records if r.get('prediction_result') == 'Positive')
    negatives = total - positives
    
    records = []
    for r in user_records:
        symptoms_list = []
        if r.get('fever'): symptoms_list.append('Fever')
        if r.get('cough'): symptoms_list.append('Cough')
        if r.get('sore_throat'): symptoms_list.append('Sore Throat')
        if r.get('shortness_of_breath'): symptoms_list.append('Shortness of Breath')
        if r.get('headache'): symptoms_list.append('Headache')
        
        prob = r['prediction_probability']
        if r['prediction_result'] == 'Positive':
            risk_class = 'risk-high'
        elif prob >= 0.25:
            risk_class = 'risk-medium'
        else:
            risk_class = 'risk-low'
            
        records.append({
            'id': r['id'],
            'name': r['name'],
            'age': r['age'],
            'gender': r['gender'],
            'symptoms': ', '.join(symptoms_list) if symptoms_list else 'None',
            'contact': 'Yes' if r.get('contact') == 1 else 'No',
            'result': r['prediction_result'],
            'probability': round(prob * 100, 1),
            'risk_class': risk_class,
            'timestamp': r['timestamp']
        })
        
    return render_template('dashboard.html', records=records, total=total, positives=positives, negatives=negatives)

@app.route('/screen')
def screen():
    if not session.get('user_id'):
        return redirect(url_for('login'))
        
    try:
        user = db_get_user_by_id(session['user_id'])
    except Exception as e:
        print(f"Error fetching user details: {e}")
        user = None
        
    return render_template('index.html', user=user)

# ==========================================================================
# USER AUTHENTICATION ROUTES
# ==========================================================================

@app.route('/register', methods=['GET', 'POST'])
def register():
    if session.get('user_id'):
        return redirect(url_for('index'))
        
    error = None
    if request.method == 'POST':
        username = request.form.get('username', '').strip().lower()
        password = request.form.get('password', '')
        name = request.form.get('name', '').strip()
        age = request.form.get('age')
        gender = request.form.get('gender')
        
        if not (username and password and name and age and gender):
            error = 'All fields are required.'
        elif not is_db_connected():
            error = 'Registration failed: Database not connected.'
        else:
            try:
                existing_user = db_get_user_by_username(username)
                if existing_user:
                    error = 'Username already registered. Please choose another one.'
                else:
                    hashed_pw = generate_password_hash(password)
                    user_id = get_next_sequence_value('users')
                    db_create_user(user_id, {
                        'id': user_id,
                        'username': username,
                        'password': hashed_pw,
                        'name': name,
                        'age': int(age),
                        'gender': gender,
                        'created_at': datetime.utcnow()
                    })
                    return redirect(url_for('login', registered=True))
            except Exception as e:
                error = f'Registration failed: {str(e)}'
                
    return render_template('register.html', error=error)

@app.route('/login', methods=['GET', 'POST'])
def login():
    if session.get('user_id'):
        return redirect(url_for('index'))
        
    error = None
    success = None
    
    if request.args.get('registered'):
        success = 'Registration successful! Please log in with your credentials.'
        
    if request.method == 'POST':
        username = request.form.get('username', '').strip().lower()
        password = request.form.get('password', '')
        
        if not (username and password):
            error = 'Please enter both username and password.'
        elif not is_db_connected():
            error = 'Login failed: Database not connected.'
        else:
            user = db_get_user_by_username(username)
            if user:
                if check_password_hash(user['password'], password):
                    session['user_id'] = user['id']
                    session['username'] = user['username']
                    session['name'] = user['name']
                    return redirect(url_for('index'))
            
            error = 'Invalid username or password.'
                
    return render_template('login.html', error=error, success=success)

@app.route('/logout')
def logout():
    session.pop('user_id', None)
    session.pop('username', None)
    session.pop('name', None)
    return redirect(url_for('login'))

# ==========================================================================
# PREDICTION INFERENCE API
# ==========================================================================

@app.route('/predict', methods=['POST'])
def predict():
    if not session.get('user_id'):
        return jsonify({'error': 'Unauthorized. Please login.'}), 401
        
    if not model:
        return jsonify({'error': 'Prediction model not loaded on server.'}), 500
        
    try:
        data = request.get_json()
        if not data:
            return jsonify({'error': 'No data provided.'}), 400
            
        name = data.get('name', 'Anonymous')
        age = int(data.get('age', 30))
        gender = data.get('gender', 'Female')
        
        # Symptoms (0 or 1)
        fever = 1 if data.get('fever') else 0
        cough = 1 if data.get('cough') else 0
        sore_throat = 1 if data.get('sore_throat') else 0
        shortness_of_breath = 1 if data.get('shortness_of_breath') else 0
        headache = 1 if data.get('headache') else 0
        contact = 1 if data.get('contact') else 0
        
        # Derived Feature
        age_60_and_above = 1 if age >= 60 else 0
        
        # Encoded Gender (Female: 0, Male: 1)
        gender_encoded = 1 if gender.lower() == 'male' else 0
        
        # Features in order: ['fever', 'cough', 'sore_throat', 'shortness_of_breath', 'headache', 'age_60_and_above', 'gender', 'contact']
        features = [[fever, cough, sore_throat, shortness_of_breath, headache, age_60_and_above, gender_encoded, contact]]
        
        # Make Prediction
        probabilities = model.predict_proba(features)[0]
        prob_positive = probabilities[1]
        
        # Decision Boundary
        result = 'Positive' if prob_positive >= 0.5 else 'Negative'
        prob_percentage = round(prob_positive * 100, 1)
        
        # Store in Database
        if not is_db_connected():
            raise Exception("Database not connected.")
            
        pred_id = get_next_sequence_value('predictions')
        db_create_prediction(pred_id, {
            'id': pred_id,
            'name': name,
            'age': age,
            'gender': gender,
            'fever': fever,
            'cough': cough,
            'sore_throat': sore_throat,
            'shortness_of_breath': shortness_of_breath,
            'headache': headache,
            'age_60_and_above': age_60_and_above,
            'contact': contact,
            'prediction_result': result,
            'prediction_probability': prob_positive,
            'timestamp': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'user_id': session['user_id']
        })
        
        # Recommendations
        recommendations = []
        if result == 'Positive':
            recommendations = [
                "Self-isolate immediately for at least 10 days or as per local guidelines.",
                "Schedule a confirmatory RT-PCR lab test at the earliest.",
                "Monitor oxygen levels regularly and consult a healthcare provider.",
                "Inform anyone you have been in close contact with recently."
            ]
        else:
            if prob_percentage > 25.0:
                recommendations = [
                    "You have a moderate risk level. Monitor your health symptoms closely.",
                    "Wear masks in public settings and practice physical distancing.",
                    "If symptoms persist or worsen, consult a doctor."
                ]
            else:
                recommendations = [
                    "Low risk level. Continue following standard health protocols.",
                    "Wash hands frequently with soap or use sanitizers.",
                    "Keep your surroundings clean and well-ventilated."
                ]
                
        return jsonify({
            'success': True,
            'name': name,
            'result': result,
            'probability': prob_percentage,
            'recommendations': recommendations
        })
        
    except Exception as e:
        return jsonify({'error': f'Failed to process prediction: {str(e)}'}), 500

# ==========================================================================
# REPORT DOWNLOAD ENGINES (WORD & PDF)
# ==========================================================================

@app.route('/download/docx/<int:pred_id>')
def download_docx(pred_id):
    if not (session.get('user_id') or session.get('admin_logged_in')):
        return redirect(url_for('login'))
        
    if not is_db_connected():
        return "Database not connected.", 500
        
    record = db_get_prediction_by_id(pred_id)
    if not record:
        return "Record not found.", 404
        
    # Security check: User can only download their own reports (unless admin is logged in)
    if not session.get('admin_logged_in') and record['user_id'] != session.get('user_id'):
        return "Access denied.", 403
        
    try:
        doc = Document()
        
        # Header / Title styling
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run("COVID-19 DIAGNOSTIC ASSESSMENT REPORT")
        title_run.bold = True
        title_run.font.size = Pt(16)
        title_run.font.name = 'Arial'
        
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub.add_run("Generated by COVID-19 AI Clinical Diagnostics Engine")
        sub_run.italic = True
        sub_run.font.size = Pt(9)
        sub_run.font.color.rgb = RGBColor(100, 116, 139)
        
        doc.add_paragraph("\n")
        
        # Clinical Details Table
        table = doc.add_table(rows=6, cols=2)
        table.style = 'Table Grid'
        
        data = [
            ("Report Reference ID", f"COVID-REPORT-00{record['id']}"),
            ("Assessment Timestamp", str(record['timestamp'])),
            ("Patient Name", str(record['name'])),
            ("Patient Age / Gender", f"{record['age']} y/o / {record['gender']}"),
            ("Diagnostic Outcome", str(record['prediction_result'])),
            ("AI Infection Probability", f"{round(record['prediction_probability'] * 100, 1)}%")
        ]
        
        for i, (k, v) in enumerate(data):
            row = table.rows[i]
            row.cells[0].text = k
            row.cells[1].text = v
            # Formatting
            row.cells[0].paragraphs[0].runs[0].font.bold = True
            row.cells[0].paragraphs[0].runs[0].font.size = Pt(10.5)
            row.cells[1].paragraphs[0].runs[0].font.size = Pt(10.5)
            if k == "Diagnostic Outcome":
                run = row.cells[1].paragraphs[0].runs[0]
                run.font.bold = True
                if v == "Positive":
                    run.font.color.rgb = RGBColor(239, 68, 68) # Red
                else:
                    run.font.color.rgb = RGBColor(16, 185, 129) # Green
            if k == "AI Infection Probability":
                row.cells[1].paragraphs[0].runs[0].font.bold = True
                
        doc.add_paragraph("\n")
        
        # Symptoms bullet grid
        doc.add_heading("Logged Symptoms & Exposure Profile:", level=2)
        
        symptoms = []
        if record['fever']: symptoms.append("Fever (Temperature >= 38 C)")
        if record['cough']: symptoms.append("Dry Cough")
        if record['sore_throat']: symptoms.append("Sore Throat")
        if record['shortness_of_breath']: symptoms.append("Shortness of Breath")
        if record['headache']: symptoms.append("Severe Headache")
        
        if symptoms:
            for symp in symptoms:
                doc.add_paragraph(symp, style='List Bullet')
        else:
            doc.add_paragraph("No symptoms logged.", style='List Bullet')
            
        exposure = "Yes" if record['contact'] else "No"
        p = doc.add_paragraph()
        run_exp_label = p.add_run("Contact with Confirmed Case: ")
        run_exp_label.bold = True
        run_exp_val = p.add_run(exposure)
        if record['contact']:
            run_exp_val.font.color.rgb = RGBColor(245, 158, 11) # Yellow
            run_exp_val.bold = True
            
        doc.add_paragraph("\n")
        
        # Medical advice list
        doc.add_heading("Clinical Guidance & Recommendations:", level=2)
        if record['prediction_result'] == 'Positive':
            recs = [
                "Self-isolate immediately for at least 10 days or as per local guidelines.",
                "Schedule a confirmatory RT-PCR lab test at the earliest.",
                "Monitor oxygen levels regularly and consult a healthcare provider.",
                "Inform anyone you have been in close contact with recently."
            ]
        else:
            prob = record['prediction_probability']
            if prob >= 0.25:
                recs = [
                    "You have a moderate risk level. Monitor your health symptoms closely.",
                    "Wear masks in public settings and practice physical distancing.",
                    "If symptoms persist or worsen, consult a doctor."
                ]
            else:
                recs = [
                    "Low risk level. Continue following standard health protocols.",
                    "Wash hands frequently with soap or use sanitizers.",
                    "Keep your surroundings clean and well-ventilated."
                ]
                
        for rec in recs:
            doc.add_paragraph(rec, style='List Bullet')
            
        doc.add_paragraph("\n\n")
        disclaimer = doc.add_paragraph()
        disclaimer_run = disclaimer.add_run(
            "Disclaimer: This report was compiled by an artificial intelligence predictive classifier using symptom records. "
            "It does not represent a lab-certified PCR diagnosis. Please consult a medical practitioner if symptoms persist."
        )
        disclaimer_run.font.size = Pt(8)
        disclaimer_run.italic = True
        disclaimer_run.font.color.rgb = RGBColor(148, 163, 184)
        
        filename = f"covid_report_{record['id']}_{datetime.now().strftime('%Y%m%d')}.docx"
        filepath = os.path.join(EXPORTS_DIR, filename)
        doc.save(filepath)
        
        return send_file(filepath, as_attachment=True, download_name=filename)
    except Exception as e:
        return f"Error generating DOCX document: {str(e)}", 500

@app.route('/download/pdf/<int:pred_id>')
def download_pdf(pred_id):
    if not (session.get('user_id') or session.get('admin_logged_in')):
        return redirect(url_for('login'))
        
    if not is_db_connected():
        return "Database not connected.", 500
        
    record = db_get_prediction_by_id(pred_id)
    if not record:
        return "Record not found.", 404
        
    # Security check: User can only download their own reports (unless admin is logged in)
    if not session.get('admin_logged_in') and record['user_id'] != session.get('user_id'):
        return "Access denied.", 403
        
    try:
        filename = f"covid_report_{record['id']}_{datetime.now().strftime('%Y%m%d')}.pdf"
        filepath = os.path.join(EXPORTS_DIR, filename)
        
        doc = SimpleDocTemplate(filepath, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
        story = []
        
        styles = getSampleStyleSheet()
        
        # Styles
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#0f172a'),
            alignment=1,
            spaceAfter=10
        )
        subtitle_style = ParagraphStyle(
            'DocSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=9,
            textColor=colors.HexColor('#64748b'),
            alignment=1,
            spaceAfter=25
        )
        section_style = ParagraphStyle(
            'SecTitle',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=13,
            leading=16,
            textColor=colors.HexColor('#1e293b'),
            spaceBefore=15,
            spaceAfter=8
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#334155')
        )
        disclaimer_style = ParagraphStyle(
            'Disclaimer',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=8,
            leading=11,
            textColor=colors.HexColor('#94a3b8'),
            spaceBefore=25
        )
        
        story.append(Paragraph("COVID-19 CLINICAL ASSESSMENT REPORT", title_style))
        story.append(Paragraph("Generated by COVID-19 AI Diagnostics Engine", subtitle_style))
        
        # Clinical parameters table
        outcome_color = '#ef4444' if record['prediction_result'] == 'Positive' else '#10b981'
        
        data = [
            [Paragraph("<b>Report ID:</b>", body_style), Paragraph(f"COVID-REPORT-00{record['id']}", body_style)],
            [Paragraph("<b>Log Date:</b>", body_style), Paragraph(str(record['timestamp']), body_style)],
            [Paragraph("<b>Patient Name:</b>", body_style), Paragraph(str(record['name']), body_style)],
            [Paragraph("<b>Age / Gender:</b>", body_style), Paragraph(f"{record['age']} y/o / {record['gender']}", body_style)],
            [Paragraph("<b>Diagnostic Status:</b>", body_style), Paragraph(f"<font color='{outcome_color}'><b>{record['prediction_result']}</b></font>", body_style)],
            [Paragraph("<b>AI Probability:</b>", body_style), Paragraph(f"<b>{round(record['prediction_probability'] * 100, 1)}%</b>", body_style)]
        ]
        
        t = Table(data, colWidths=[180, 320])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('LEFTPADDING', (0, 0), (-1, -1), 10),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#cbd5e1')),
            ('LINEBELOW', (0, 0), (-1, -2), 0.5, colors.HexColor('#e2e8f0')),
        ]))
        story.append(t)
        story.append(Spacer(1, 10))
        
        # Symptoms list
        story.append(Paragraph("Reported Symptoms & Exposure", section_style))
        
        symptoms = []
        if record['fever']: symptoms.append("Fever")
        if record['cough']: symptoms.append("Dry Cough")
        if record['sore_throat']: symptoms.append("Sore Throat")
        if record['shortness_of_breath']: symptoms.append("Shortness of Breath")
        if record['headache']: symptoms.append("Severe Headache")
        
        symp_text = ", ".join(symptoms) if symptoms else "No clinical symptoms reported."
        story.append(Paragraph(f"<b>Symptom Checklist:</b> {symp_text}", body_style))
        story.append(Spacer(1, 4))
        
        exposure = "Yes" if record['contact'] else "No"
        story.append(Paragraph(f"<b>Recent Contact with Confirmed COVID Patient:</b> {exposure}", body_style))
        story.append(Spacer(1, 10))
        
        # Advice list
        story.append(Paragraph("Clinical Advice & Directives", section_style))
        if record['prediction_result'] == 'Positive':
            recs = [
                "Self-isolate immediately for at least 10 days or as per local guidelines.",
                "Schedule a confirmatory RT-PCR lab test at the earliest.",
                "Monitor oxygen levels regularly and consult a healthcare provider.",
                "Inform anyone you have been in close contact with recently."
            ]
        else:
            prob = record['prediction_probability']
            if prob >= 0.25:
                recs = [
                    "You have a moderate risk level. Monitor your health symptoms closely.",
                    "Wear masks in public settings and practice physical distancing.",
                    "If symptoms persist or worsen, consult a doctor."
                ]
            else:
                recs = [
                    "Low risk level. Continue following standard health protocols.",
                    "Wash hands frequently with soap or use sanitizers.",
                    "Keep your surroundings clean and well-ventilated."
                ]
                
        for rec in recs:
            story.append(Paragraph(f"&bull; {rec}", body_style))
            story.append(Spacer(1, 4))
            
        story.append(Spacer(1, 15))
        story.append(Paragraph("<i>Disclaimer: This report was compiled by an artificial intelligence predictive classifier using symptom records. It does not represent a lab-certified PCR diagnosis. Please consult a medical practitioner if symptoms persist.</i>", disclaimer_style))
        
        doc.build(story)
        return send_file(filepath, as_attachment=True, download_name=filename)
    except Exception as e:
        return f"Error generating PDF document: {str(e)}", 500

# ==========================================================================
# ADMINISTRATIVE CONTROL PANEL ROUTES
# ==========================================================================

@app.route('/admin/login', methods=['GET', 'POST'])
def admin_login():
    if session.get('admin_logged_in'):
        return redirect(url_for('admin_dashboard'))
        
    error = None
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        
        if username == ADMIN_USERNAME and password == ADMIN_PASSWORD:
            session['admin_logged_in'] = True
            return redirect(url_for('admin_dashboard'))
        else:
            error = 'Invalid credentials. Please try again.'
            
    return render_template('admin_login.html', error=error)

@app.route('/admin/dashboard')
def admin_dashboard():
    if not session.get('admin_logged_in'):
        return redirect(url_for('admin_login'))
        
    if not is_db_connected():
        records_raw = []
    else:
        try:
            records_raw = db_get_all_predictions()
            users_map = db_get_users_map()
            for r in records_raw:
                r['user_username'] = users_map.get(r.get('user_id'), 'Guest')
        except Exception as e:
            print(f"Error fetching admin dashboard records: {e}")
            records_raw = []
    
    # Calculate stats
    total_screenings = len(records_raw)
    positive_cases = sum(1 for r in records_raw if r.get('prediction_result') == 'Positive')
    negative_cases = total_screenings - positive_cases
    positivity_rate = round((positive_cases / total_screenings * 100), 1) if total_screenings > 0 else 0
    
    # Symptom counts
    fever_count = sum(1 for r in records_raw if r.get('fever') == 1)
    cough_count = sum(1 for r in records_raw if r.get('cough') == 1)
    sore_throat_count = sum(1 for r in records_raw if r.get('sore_throat') == 1)
    sob_count = sum(1 for r in records_raw if r.get('shortness_of_breath') == 1)
    headache_count = sum(1 for r in records_raw if r.get('headache') == 1)
    contact_count = sum(1 for r in records_raw if r.get('contact') == 1)
    
    stats = {
        'total_screenings': total_screenings,
        'positive_cases': positive_cases,
        'negative_cases': negative_cases,
        'positivity_rate': positivity_rate,
        'fever': fever_count,
        'cough': cough_count,
        'sore_throat': sore_throat_count,
        'shortness_of_breath': sob_count,
        'headache': headache_count,
        'contact': contact_count
    }
    
    # Format records for rendering
    records = []
    for r in records_raw:
        symptoms_list = []
        if r.get('fever'): symptoms_list.append('Fever')
        if r.get('cough'): symptoms_list.append('Cough')
        if r.get('sore_throat'): symptoms_list.append('Sore Throat')
        if r.get('shortness_of_breath'): symptoms_list.append('Shortness of Breath')
        if r.get('headache'): symptoms_list.append('Headache')
        
        # Risk color
        prob = r['prediction_probability']
        if r['prediction_result'] == 'Positive':
            risk_class = 'risk-high'
        elif prob >= 0.25:
            risk_class = 'risk-medium'
        else:
            risk_class = 'risk-low'
            
        records.append({
            'id': r['id'],
            'name': r['name'],
            'age': r['age'],
            'gender': r['gender'],
            'symptoms': ', '.join(symptoms_list) if symptoms_list else 'None',
            'contact': 'Yes' if r.get('contact') == 1 else 'No',
            'result': r['prediction_result'],
            'probability': round(prob * 100, 1),
            'risk_class': risk_class,
            'timestamp': r['timestamp'],
            'user_username': r.get('user_username') or 'Guest'
        })
        
    return render_template('admin_dashboard.html', stats=stats, records=records)

@app.route('/admin/delete/<int:record_id>', methods=['POST'])
def delete_record(record_id):
    if not session.get('admin_logged_in'):
        return jsonify({'error': 'Unauthorized'}), 401
        
    if not is_db_connected():
        return jsonify({'error': 'Database not connected.'}), 500
        
    try:
        db_delete_prediction(record_id)
        return jsonify({'success': True, 'message': 'Record deleted successfully.'})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/admin/export')
def export_csv():
    if not session.get('admin_logged_in'):
        return redirect(url_for('admin_login'))
        
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f"predictions_export_{timestamp}.csv"
    export_path = os.path.join(EXPORTS_DIR, filename)
    
    if not is_db_connected():
        return "Database not connected.", 500
        
    try:
        records_raw = db_get_all_predictions()
        users_map = db_get_users_map()
        for r in records_raw:
            r['user_username'] = users_map.get(r.get('user_id'), 'Guest')
        
        # Define headers
        headers = [
            'id', 'name', 'age', 'gender', 'fever', 'cough', 
            'sore_throat', 'shortness_of_breath', 'headache', 
            'age_60_and_above', 'contact', 'prediction_result', 
            'prediction_probability', 'timestamp', 'user_id', 'user_username'
        ]
        
        with open(export_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for row in records_raw:
                # Compile row values in order of headers
                row_data = [row.get(h, '') for h in headers]
                writer.writerow(row_data)
                
        return send_file(export_path, as_attachment=True, download_name=filename)
    except Exception as e:
        return f"Error exporting database: {str(e)}", 500

@app.route('/admin/logout')
def admin_logout():
    session.pop('admin_logged_in', None)
    return redirect(url_for('admin_login'))

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
